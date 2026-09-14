#!/usr/bin/env python3
"""把结构化的专利镜头数据（spec.json）生成 Excel 工作簿。

    python3 build_workbook.py spec.json -o 输出.xlsx

spec.json 的字段说明见 references/spec-schema.md。
"""
import json, sys, os, argparse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from glasslib import load, match, grade, tolerances, exact_vendor

FONT = 'Arial'
HDR  = PatternFill('solid', fgColor='1F4E79')
SUB  = PatternFill('solid', fgColor='D9D9D9')
FILL = ['E2EFDA', 'FFF2CC', 'DDEBF7']            # 各厂商列的底色
ORIG = PatternFill('solid', fgColor='F0E6F6')
STOP = PatternFill('solid', fgColor='FCE4D6')
ALT  = PatternFill('solid', fgColor='F2F2F2')
ASPH = PatternFill('solid', fgColor='EDE7F6')
WARN = PatternFill('solid', fgColor='FCE4D6')
_t = Side('thin', color='B0B0B0')
BOX = Border(_t, _t, _t, _t)


def _f(sz=10, bold=False, color='000000', italic=False):
    return Font(FONT, size=sz, bold=bold, color=color, italic=italic)


def _hdr(ws, row, labels, widths=None):
    for c, h in enumerate(labels, 1):
        x = ws.cell(row, c, h)
        x.font = _f(10, True, 'FFFFFF'); x.fill = HDR; x.border = BOX
        x.alignment = Alignment('center', 'center', wrap_text=True)
    ws.row_dimensions[row].height = 32
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w


def _num(x, c, kind):
    fmt = {'r': '0.0000', 'd': '0.0000', 'nd': '0.00000', 'vd': '0.00',
           'dnd': '+0.00000;-0.00000;0', 'dvd': '+0.00;-0.00;0', 'g': '0.0000'}
    if kind in fmt:
        c.number_format = fmt[kind]


def resolve_d(s, emb):
    """D 可以是数字，也可以是变量名（如 'D1'），后者按各对焦/变倍状态展开。"""
    d = s.get('D')
    states = emb.get('states') or []
    var = emb.get('variable') or {}
    if isinstance(d, str) and d in var:
        return [var[d].get(st) for st in states], d
    return [d] * max(len(states), 1), None


def build(spec, out):
    vendors = [v.upper() for v in spec.get('vendors', ['HOYA', 'OHARA'])]
    embs = spec['embodiments']
    nds = [s['nd'] for e in embs for s in e['surfaces'] if s.get('nd')]
    # 容差默认按 nd 的小数位自动判定；float 无法保留末尾 0，若专利明明给了 5 位
    # 而数值恰好都以 0 结尾，可用 spec 的 nd_decimals 或 tolerance 显式指定。
    if spec.get('tolerance'):
        tol = tuple(spec['tolerance'])
    elif spec.get('nd_decimals'):
        d = int(spec['nd_decimals'])
        tol = (0.0005, 0.05) if d >= 4 else ((0.002, 0.15) if d == 3 else (0.006, 0.25))
    else:
        tol = tolerances(nds)
    libs = {v: load(v) for v in vendors}
    wb = Workbook()

    # ---------------- 主表：全实施例镜片数据 ----------------
    ws = wb.active; ws.title = '镜片数据_全实施例'
    ws['A1'] = spec.get('title', '专利镜头数据'); ws['A1'].font = _f(13, True)
    ws['A2'] = spec.get('subtitle', ''); ws['A2'].font = _f(9, italic=True, color='808080')
    if spec.get('note'):
        ws['A3'] = spec['note']; ws['A3'].font = _f(9, color='C00000')

    extra_cols = []
    for e in embs:
        for k in e.get('extra_columns', []):
            if k not in extra_cols:
                extra_cols.append(k)
    multi = max(len(e.get('states') or []) for e in embs) > 1
    states0 = max((e.get('states') or [] for e in embs), key=len)

    base = ['实施例', '面号', '面型', '透镜', '组元', 'R 曲率半径']
    if multi:
        base += ['D (印刷)'] + ['D @ %s' % s for s in states0]
    else:
        base += ['D 间隔']
    base += extra_cols + ['Nd', 'Vd']
    per_v = ['{v} 牌号', '{v} nd', '{v} vd', 'Δnd', 'Δvd', '评级', '{v} 备选']
    head = list(base)
    for v in vendors:
        head += [c.format(v=v) for c in per_v]
    head += ['原设计推定', '备注']
    NB = len(base)
    widths = ([14, 7, 8, 7, 16, 13] + ([12] + [11] * len(states0) if multi else [11])
              + [10] * len(extra_cols) + [10, 8]
              + [14, 10, 8, 10, 8, 12, 30] * len(vendors) + [24, 40])
    HR = 5
    _hdr(ws, HR, head, widths)

    r = HR + 1
    for ei, e in enumerate(embs):
        states = e.get('states') or []
        for si, s in enumerate(e['surfaces']):
            i = s.get('i', si + 1)
            is_stop = str(i).upper() in ('STO', 'STOP', '光阑', '絞り') or s.get('stop')
            asph = (s.get('type') or '').startswith(('非球', 'ASP', 'asph'))
            ws.cell(r, 1, e['name'] if si == 0 else '')
            ws.cell(r, 2, i)
            ws.cell(r, 3, s.get('type', '球面'))
            ws.cell(r, 4, s.get('lens', ''))
            ws.cell(r, 5, s.get('group', ''))
            R = s.get('R')
            ws.cell(r, 6, 'INFINITY' if R in (None, 'inf', 'INF', 'Infinity', '无穷大', '∞') else R)
            col = 7
            dvals, dname = resolve_d(s, e)
            if multi:
                ws.cell(r, col, dname or s.get('D')); col += 1
                for k in range(len(states0)):
                    ws.cell(r, col, dvals[k] if k < len(dvals) else None); col += 1
            else:
                ws.cell(r, col, s.get('D')); col += 1
            for k in extra_cols:
                ws.cell(r, col, (s.get('extra') or {}).get(k)); col += 1
            nd, vd = s.get('nd'), s.get('vd')
            ws.cell(r, col, nd); ws.cell(r, col + 1, vd)
            ndc, vdc = col, col + 1
            col += 2
            if nd:
                for vi, v in enumerate(vendors):
                    top = match(nd, vd, libs[v], tol, aspheric=asph)
                    g = top[0]
                    ws.cell(r, col,     g['name'])
                    ws.cell(r, col + 1, round(g['nd'], 5))
                    ws.cell(r, col + 2, round(g['vd'], 2))
                    ws.cell(r, col + 3, '=%s%d-%s%d' % (get_column_letter(col + 1), r, get_column_letter(ndc), r))
                    ws.cell(r, col + 4, '=%s%d-%s%d' % (get_column_letter(col + 2), r, get_column_letter(vdc), r))
                    ws.cell(r, col + 5, grade(g, nd, vd, tol))
                    ws.cell(r, col + 6, ' / '.join('%s(%.5f/%.2f)' % (x['name'], x['nd'], x['vd']) for x in top[1:]))
                    col += 7
                ws.cell(r, col, exact_vendor(nd, vd, list(libs.values()), tol))
            col = NB + 7 * len(vendors) + 1
            ws.cell(r, col + 1, s.get('note', ''))
            # 排版
            for c in range(1, len(head) + 1):
                x = ws.cell(r, c); x.font = _f(10); x.border = BOX
                x.alignment = Alignment(
                    'left' if c >= NB + 1 and (c - NB) % 7 == 0 or c in (5, len(head)) else
                    ('center' if c in (1, 2, 3, 4) else 'right'), 'center')
                if c == 6: _num(None, x, 'r')
                if 7 <= c < NB - 1: _num(None, x, 'd')
                if c == ndc: _num(None, x, 'nd')
                if c == vdc: _num(None, x, 'vd')
                for vi in range(len(vendors)):
                    b = NB + 7 * vi
                    if c == b + 2: _num(None, x, 'nd')
                    if c == b + 3: _num(None, x, 'vd')
                    if c == b + 4: _num(None, x, 'dnd')
                    if c == b + 5: _num(None, x, 'dvd')
                    if b + 1 <= c <= b + 7:
                        x.fill = PatternFill('solid', fgColor=FILL[vi % len(FILL)])
                if c == NB + 7 * len(vendors) + 1: x.fill = ORIG
                if c <= NB:
                    if is_stop: x.fill = STOP
                    elif asph:  x.fill = ASPH
                    elif ei % 2: x.fill = ALT
                if c == len(head) and str(s.get('note', '')).startswith('★'):
                    x.fill = WARN
            if dname:
                for k in range(len(states0)):
                    ws.cell(r, 8 + k).font = _f(10, True, 'C00000')
            r += 1
        r += 1
    ws.freeze_panes = ws.cell(HR + 1, 6).coordinate

    # ---------------- 各种数据 ----------------
    if any(e.get('general') or e.get('variable') or e.get('aspheric') for e in embs):
        g = wb.create_sheet('各种数据')
        g['A1'] = '各实施例的规格参数 / 可变间隔 / 非球面系数'; g['A1'].font = _f(13, True)
        row = 3
        keys = []
        for e in embs:
            for k, _ in (e.get('general') or []):
                if k not in keys: keys.append(k)
        if keys:
            _hdr(g, row, ['项目'] + [e['name'] for e in embs], [30] + [15] * len(embs))
            for j, k in enumerate(keys):
                g.cell(row + 1 + j, 1, k).font = _f(10)
                g.cell(row + 1 + j, 1).border = BOX
                for c, e in enumerate(embs, 2):
                    val = dict(e.get('general') or []).get(k)
                    x = g.cell(row + 1 + j, c, val); x.font = _f(10); x.border = BOX
                    x.alignment = Alignment('center', 'center')
                    if isinstance(val, float): x.number_format = '0.00'
            row += len(keys) + 3
        for e in embs:
            if e.get('variable'):
                g.cell(row, 1, '%s — 可变间隔' % e['name']).font = _f(11, True, '1F4E79')
                _hdr(g, row + 1, ['间隔'] + list(e.get('states') or []))
                for j, (k, d) in enumerate(e['variable'].items()):
                    g.cell(row + 2 + j, 1, k).font = _f(10)
                    g.cell(row + 2 + j, 1).border = BOX
                    for c, st in enumerate(e.get('states') or [], 2):
                        x = g.cell(row + 2 + j, c, d.get(st)); x.font = _f(10)
                        x.border = BOX; x.number_format = '0.0000'
                        x.alignment = Alignment('center', 'center')
                row += len(e['variable']) + 4
        for e in embs:
            if e.get('aspheric'):
                g.cell(row, 1, '%s — 非球面系数  z = ch²/(1+√(1-(1+K)c²h²)) + ΣAi·h^i' % e['name']).font = _f(11, True, '1F4E79')
                cols = ['面'] + [k for k in e['aspheric'][0].keys() if k != 'surface']
                _hdr(g, row + 1, cols)
                for j, a in enumerate(e['aspheric']):
                    x = g.cell(row + 2 + j, 1, a.get('surface')); x.font = _f(10); x.border = BOX
                    x.alignment = Alignment('center', 'center')
                    for c, k in enumerate(cols[1:], 2):
                        x = g.cell(row + 2 + j, c, a.get(k)); x.font = _f(10); x.border = BOX
                        x.number_format = '0.0000E+00'
                        x.alignment = Alignment('center', 'center')
                row += len(e['aspheric']) + 4
        g.column_dimensions['A'].width = 34

    # ---------------- 玻璃汇总 ----------------
    sm = wb.create_sheet('玻璃汇总')
    sm['A1'] = '全部材料（去重）— %s 对应表' % ' / '.join(vendors); sm['A1'].font = _f(13, True)
    head = ['Nd', 'Vd', '使用位置']
    for v in vendors:
        head += ['%s 牌号' % v, 'nd', 'vd', 'Δnd', 'Δvd', '评级', '备选']
    head += ['原设计推定']
    _hdr(sm, 3, head, [10, 8, 32] + [15, 10, 8, 10, 8, 12, 32] * len(vendors) + [24])
    use = {}
    for e in embs:
        for s in e['surfaces']:
            if s.get('nd'):
                use.setdefault((s['nd'], s['vd']), []).append('%s·%s' % (e['name'], s.get('lens') or s.get('i')))
    rr = 4
    for (nd, vd), loc in sorted(use.items()):
        sm.cell(rr, 1, nd); sm.cell(rr, 2, vd); sm.cell(rr, 3, '，'.join(loc))
        c = 4
        for vi, v in enumerate(vendors):
            top = match(nd, vd, libs[v], tol); g0 = top[0]
            sm.cell(rr, c, g0['name']); sm.cell(rr, c + 1, round(g0['nd'], 5)); sm.cell(rr, c + 2, round(g0['vd'], 2))
            sm.cell(rr, c + 3, '=%s%d-A%d' % (get_column_letter(c + 1), rr, rr))
            sm.cell(rr, c + 4, '=%s%d-B%d' % (get_column_letter(c + 2), rr, rr))
            sm.cell(rr, c + 5, grade(g0, nd, vd, tol))
            sm.cell(rr, c + 6, ' / '.join('%s(%.5f/%.2f)' % (x['name'], x['nd'], x['vd']) for x in top[1:]))
            c += 7
        sm.cell(rr, c, exact_vendor(nd, vd, list(libs.values()), tol))
        for cc in range(1, len(head) + 1):
            x = sm.cell(rr, cc); x.font = _f(10); x.border = BOX
            x.alignment = Alignment('left' if cc in (3, len(head)) or (cc > 3 and (cc - 3) % 7 == 0) else 'center', 'center')
            if cc in (1,): x.number_format = '0.00000'
            if cc in (2,): x.number_format = '0.00'
            for vi in range(len(vendors)):
                b = 3 + 7 * vi
                if cc == b + 2: x.number_format = '0.00000'
                if cc == b + 3: x.number_format = '0.00'
                if cc == b + 4: x.number_format = '+0.00000;-0.00000;0'
                if cc == b + 5: x.number_format = '+0.00;-0.00;0'
                if b + 1 <= cc <= b + 7: x.fill = PatternFill('solid', fgColor=FILL[vi % len(FILL)])
            if cc == len(head): x.fill = ORIG
        rr += 1
    sm.freeze_panes = 'C4'

    # ---------------- 粘贴用（每个实施例一张纯数据表） ----------------
    for e in embs:
        name = ('粘贴用_' + e['name'])[:31]
        w = wb.create_sheet(name)
        states = e.get('states') or []
        head = ['面号', 'R'] + (['D @ %s' % s for s in states] if len(states) > 1 else ['D']) \
               + e.get('extra_columns', []) + ['Nd', 'Vd'] + vendors
        w.append(head)
        for c in range(1, len(head) + 1):
            x = w.cell(1, c); x.font = _f(10, True, 'FFFFFF'); x.fill = HDR
        for s in e['surfaces']:
            R = s.get('R')
            dvals, _ = resolve_d(s, e)
            row = [s.get('i'), 'INF' if R in (None, 'inf', 'INF', 'Infinity', '无穷大', '∞') else R]
            row += (dvals if len(states) > 1 else [dvals[0]])
            row += [(s.get('extra') or {}).get(k) for k in e.get('extra_columns', [])]
            row += [s.get('nd') or '', s.get('vd') or '']
            asph = (s.get('type') or '').startswith(('非球', 'ASP', 'asph'))
            row += [match(s['nd'], s['vd'], libs[v], tol, aspheric=asph)[0]['name'] if s.get('nd') else ''
                    for v in vendors]
            w.append(row)
        nD = len(states) if len(states) > 1 else 1
        for row in w.iter_rows(min_row=2):
            for x in row:
                x.font = _f(10)
                if x.column == 2 or 3 <= x.column < 3 + nD + len(e.get('extra_columns', [])):
                    x.number_format = '0.0000'
                if x.column == 3 + nD + len(e.get('extra_columns', [])): x.number_format = '0.00000'
                if x.column == 4 + nD + len(e.get('extra_columns', [])): x.number_format = '0.00'
        for c in range(1, len(head) + 1):
            w.column_dimensions[get_column_letter(c)].width = 14 if c > len(head) - len(vendors) else 11
        w.freeze_panes = 'A2'

    wb.save(out)
    return out, tol


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('spec')
    ap.add_argument('-o', '--out', default=None)
    a = ap.parse_args()
    spec = json.load(open(a.spec, encoding='utf-8'))
    out = a.out or (spec.get('patent', 'patent').replace(' ', '_') + '_镜片数据.xlsx')
    path, tol = build(spec, out)
    print('已生成: %s   (匹配容差 nd±%.4f / vd±%.2f)' % (path, tol[0], tol[1]))


if __name__ == '__main__':
    main()
